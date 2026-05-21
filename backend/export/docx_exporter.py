from __future__ import annotations

from io import BytesIO

try:
	from docx import Document
except Exception:  # pragma: no cover - optional dependency fallback
	Document = None


class DocxExporter:
	def export(self, title: str, content: str) -> bytes:
		if Document is None:
			return f"{title}\n\n{content}\n".encode("utf-8")
		document = Document()
		document.add_heading(title, level=1)
		for paragraph in content.splitlines():
			document.add_paragraph(paragraph)
		buffer = BytesIO()
		document.save(buffer)
		return buffer.getvalue()
